"""The SIPRI Milex workbook into Bronze records, for Indonesia.

The generic workbook reader cannot serve this file. It yields one row per
spreadsheet row keyed by column letter, which for Milex means 199 countries ×
7 sheets of cells whose meaning lives in a header row four rows above and a
region heading somewhere above that. Nothing downstream can turn that back
into `Indonesia, 2025, 15025.4 US$ m.`.

So this reader knows the shape: each sheet is one measure, its header row is
the one starting `Country`, and every column header that is a year is a
period. It keeps Indonesia and drops the rest — the warehouse is Indonesian
(program.md §1), and landing the whole workbook means the other two hundred
countries are still in RAW if that ever changes.

Three things are read that a cell value alone does not carry:

- **The measure.** One sheet is US$ at constant prices, the next the same
  figures as a share of GDP. They are different series and must not land on
  one indicator, so each sheet names its own.
- **The unit.** SIPRI states it in the sheet's preamble, not in the cells. The
  two share sheets store a fraction — 0.0086, formatted as `0.86%` — so their
  unit says *share of GDP*, not percent: multiplying here would be reading a
  meaning into Bronze, and calling a fraction a percentage would be wrong by a
  hundred.
- **Whether SIPRI stands behind the figure.** Blue is a SIPRI estimate and red
  is highly uncertain data, stated in the preamble and encoded nowhere but the
  font. Read here because the font does not survive into Bronze, and a figure
  the publisher flags as uncertain should not arrive looking like a
  measurement.

`PARSER_VERSION` is not bumped for this reader. A bump makes every extractor
re-read the whole corpus, and it earns that where a reader replaces one that
already produced rows — the FRED CSVs at version 9. Nothing has ever been
extracted from this source, so its first run reads it under the current
version like any newly landed document.

Absence markers — `...` for unavailable, `xxx` for a country that did not
exist that year — yield no record. They are SIPRI saying there is no figure,
not a figure that failed to parse, and Indonesia's 1949–1956 gap would
otherwise become eight observations asserting nothing.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

from .base import ExtractionError, Extractor, Landed
from .tabular import number_text

SOURCE_SLUG = "sipri-milex"

#: The one country this warehouse keeps. Compared case-insensitively against
#: the first column, which is where SIPRI writes country names; region and
#: sub-region headings sit in the same column and match nothing.
COUNTRY = "Indonesia"

WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}

#: SIPRI's own markers for "there is no figure here", in every spelling the
#: workbook uses — the spacing differs between sheets (`...` and `. .`).
_ABSENT = {"...", "..", ".", "…", "xxx", "-", "–"}

#: Indexed colours from the legacy palette: 12 is blue (a SIPRI estimate), 10
#: is red (highly uncertain). Also matched as RGB, since a future rewrite of
#: the workbook could carry the same colours that way.
_BLUE = {"indexed:12", "FF0000FF", "0000FF"}
_RED = {"indexed:10", "FFFF0000", "FF0000"}

_BASE_YEAR = re.compile(r"\((\d{4})\)")
_YEAR = re.compile(r"^(19|20)\d{2}$")


@dataclass(frozen=True, slots=True)
class Measure:
    """One sheet of the workbook: a series, and what its numbers are in."""

    #: The Silver indicator these figures belong to. Readable rather than a
    #: code: there are six of them, and `sipri_milex_share_gdp` in a URL is
    #: worth more than eight characters of base36.
    indicator: str

    #: What a reader sees beside the identifier.
    name: str

    #: SIPRI's unit, from the sheet's preamble. `{base}` is filled with the
    #: constant-price base year, which moves every release.
    unit: str

    #: Matched against the lowercased sheet name. SIPRI renames the constant
    #: price sheet each year — `Constant (2024) US$` — so sheets are claimed
    #: by prefix rather than by their full name.
    sheet_prefix: str


#: What is read, and what is left. The regional totals sheet holds no country
#: rows; the financial-years sheet restates the local currency figures against
#: fiscal years the sheet describes in prose ("Till: 1999 Apr.–Mar."), which
#: is a period this warehouse cannot express and would silently read as a
#: calendar year.
MEASURES: tuple[Measure, ...] = (
    Measure(
        indicator="sipri_milex_constant_usd",
        name="Military expenditure, constant US$",
        unit="US$ m., constant {base} prices and exchange rates",
        sheet_prefix="constant",
    ),
    Measure(
        indicator="sipri_milex_current_usd",
        name="Military expenditure, current US$",
        unit="US$ m., current prices and exchange rates",
        sheet_prefix="current us$",
    ),
    Measure(
        indicator="sipri_milex_share_gdp",
        name="Military expenditure as a share of GDP",
        unit="share of GDP",
        sheet_prefix="share of gdp",
    ),
    Measure(
        indicator="sipri_milex_per_capita",
        name="Military expenditure per capita",
        unit="US$ per capita, current prices",
        sheet_prefix="per capita",
    ),
    Measure(
        indicator="sipri_milex_share_govt_spending",
        name="Military expenditure as a share of government spending",
        unit="share of government spending",
        sheet_prefix="share of govt",
    ),
    Measure(
        indicator="sipri_milex_local_currency",
        name="Military expenditure, local currency",
        # Overridden per row by the sheet's own Currency column, which is the
        # only place the workbook says whose money the figure is in.
        unit="local currency, current prices",
        sheet_prefix="local currency calendar",
    ),
)


def measure_for(sheet_name: str) -> Measure | None:
    """The measure a sheet holds, or None for a sheet that holds none."""
    name = sheet_name.strip().lower()
    for measure in MEASURES:
        if name.startswith(measure.sheet_prefix):
            return measure
    return None


def _is_absent(text: str) -> bool:
    """Whether a cell is one of SIPRI's ways of writing "no figure"."""
    collapsed = "".join(text.split()).lower()
    return collapsed in _ABSENT or not collapsed


def _colour(cell: Any) -> str | None:
    font = getattr(cell, "font", None)
    colour = getattr(font, "color", None)
    if colour is None:
        return None
    if colour.type == "indexed":
        return f"indexed:{colour.indexed}"
    if colour.type == "rgb" and isinstance(colour.rgb, str):
        return colour.rgb.upper()
    return None


def _load(content: bytes) -> Any:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the extra
        raise ExtractionError(
            SOURCE_SLUG, "openpyxl is not installed; add the `agencies` extra"
        ) from exc

    # Not read-only: the estimate and uncertainty flags are font colours, and
    # a read-only worksheet does not carry styles.
    return load_workbook(io.BytesIO(content), data_only=True)


def _header_row(sheet: Any, limit: int = 12) -> int | None:
    """The row whose first cell reads `Country`.

    Found rather than fixed: the preamble runs to four lines on one sheet and
    seven on another, and it grows whenever SIPRI adds a caveat.
    """
    for row in range(1, min(limit, sheet.max_row) + 1):
        if str(sheet.cell(row=row, column=1).value or "").strip().lower() == "country":
            return row
    return None


def _year_columns(sheet: Any, header: int) -> list[tuple[int, str]]:
    """Every column of the header row that is a year, in sheet order."""
    years: list[tuple[int, str]] = []
    for column in range(1, sheet.max_column + 1):
        label = number_text(sheet.cell(row=header, column=column).value)
        if _YEAR.match(label):
            years.append((column, label))
    return years


def _labelled_columns(sheet: Any, header: int) -> dict[str, int]:
    """The named columns beside the years: Notes, Currency, and the like."""
    labels: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        label = str(sheet.cell(row=header, column=column).value or "").strip().lower()
        if label and not _YEAR.match(label):
            labels.setdefault(label, column)
    return labels


def _country_rows(sheet: Any, header: int, country: str) -> list[int]:
    wanted = country.strip().lower()
    return [
        row
        for row in range(header + 1, sheet.max_row + 1)
        if str(sheet.cell(row=row, column=1).value or "").strip().lower() == wanted
    ]


class SipriMilexExtractor(Extractor):
    """One Bronze record per Indonesian figure in the Milex workbook."""

    target = "records"

    def __init__(self, country: str = COUNTRY) -> None:
        self._country = country

    def handles(self, landed: Landed) -> bool:
        return landed.source_slug == SOURCE_SLUG and landed.path.suffix.lower() in WORKBOOK_SUFFIXES

    def extract(self, landed: Landed) -> Iterator[dict[str, Any]]:
        try:
            workbook = _load(landed.path.read_bytes())
        except ExtractionError:
            raise
        except Exception as exc:  # noqa: BLE001 - surfaced with the path attached
            raise ExtractionError(str(landed.path), f"unreadable workbook: {exc}") from exc

        try:
            rows = self._rows(workbook, landed)
            number = 0
            for row in rows:
                number += 1
                yield {
                    "dataset": landed.dataset or "milex",
                    "row_number": number,
                    "columns": row,
                }
            if number == 0:
                # Not an empty file: the workbook parsed and holds no figure
                # for the one country this reader is for, which means SIPRI
                # renamed the row, renamed the sheets, or dropped the country.
                raise ExtractionError(
                    str(landed.path),
                    f"workbook holds no figures for {self._country}; "
                    "its sheets or its country names have changed",
                )
        finally:
            workbook.close()

    def _rows(self, workbook: Any, landed: Landed) -> Iterator[dict[str, str]]:
        edition = str((landed.extra or {}).get("covers_to") or "")

        for sheet_name in workbook.sheetnames:
            measure = measure_for(sheet_name)
            if measure is None:
                continue
            sheet = workbook[sheet_name]
            header = _header_row(sheet)
            if header is None:
                continue

            years = _year_columns(sheet, header)
            labelled = _labelled_columns(sheet, header)
            base = _BASE_YEAR.search(sheet_name)
            unit = measure.unit.format(base=base.group(1) if base else "")

            for row in _country_rows(sheet, header, self._country):
                currency = self._cell(sheet, row, labelled.get("currency"))
                notes = self._cell(sheet, row, labelled.get("notes"))

                for column, year in years:
                    cell = sheet.cell(row=row, column=column)
                    value = number_text(cell.value)
                    if _is_absent(value):
                        continue
                    colour = _colour(cell)

                    yield {
                        "indicator": measure.indicator,
                        "series_name": measure.name,
                        "country": self._country,
                        "year": year,
                        "value": value,
                        # The currency column where the sheet has one, since
                        # "local currency" is not a unit anyone can chart.
                        "unit": f"{currency}, current prices" if currency else unit,
                        "currency": currency,
                        "measure": sheet_name.strip(),
                        # SIPRI's footnote marks for the country, kept as
                        # written: they point into the Footnotes sheet, which
                        # says what a country's figures include.
                        "notes": notes,
                        "estimate": "1" if colour in _BLUE else "",
                        "uncertain": "1" if colour in _RED else "",
                        "edition": edition,
                        "publisher": "SIPRI",
                    }

    @staticmethod
    def _cell(sheet: Any, row: int, column: int | None) -> str:
        if column is None:
            return ""
        return number_text(sheet.cell(row=row, column=column).value)

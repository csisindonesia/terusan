"""Excel workbooks and the archives they arrive in.

Indonesian agencies publish statistics as workbooks — often zipped, often with
one sheet per table and a title row nobody promises to keep in the same place.
So extraction here is deliberately shallow: every sheet becomes rows of cells,
keyed by column letter, with the sheet name carried alongside.

That is not the finished shape. Reading a particular table out of a particular
sheet is a source's own job, because it needs to know that OJK's Tabel 1.33.a
starts three rows down and that Bank Indonesia's per-city index table is
numbered 8 rather than 6. What this gives such a reader is the cells, already in
Bronze, without a second trip to the agency.
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Iterator
from typing import Any

from .base import ExtractionError, Extractor, Landed

WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
ARCHIVE_SUFFIXES = {".zip"}

#: SpreadsheetML 2003 — an XML dialect Excel reads as a workbook, and what
#: several government portals serve when they say "Excel". It is not a ZIP, so
#: openpyxl cannot open it.
SPREADSHEETML_NS = "urn:schemas-microsoft-com:office:spreadsheet"
SPREADSHEETML_SUFFIXES = {".xml"}

#: Cells past this are almost always formatting rather than data, and reading
#: them turns one workbook into millions of empty Bronze rows.
MAX_ROWS_PER_SHEET = 20_000
MAX_COLUMNS = 80


def _load(content: bytes) -> Any:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the extra
        raise ExtractionError(
            "workbook", "openpyxl is not installed; add the `agencies` extra"
        ) from exc

    # Values, not formulae: a formula string is not a figure, and these files
    # are published as results rather than as live models.
    return load_workbook(io.BytesIO(content), data_only=True, read_only=True)


def _sheet_rows(workbook: Any, dataset: str) -> Iterator[dict[str, Any]]:
    from openpyxl.utils import get_column_letter

    for sheet in workbook.worksheets:
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if index > MAX_ROWS_PER_SHEET:
                break
            cells = {
                get_column_letter(column): "" if value is None else str(value)
                for column, value in enumerate(row[:MAX_COLUMNS], start=1)
            }
            # A row of empty cells carries nothing and there are thousands of
            # them below the last table in a typical sheet.
            if not any(cells.values()):
                continue

            yield {
                "dataset": dataset,
                "row_number": index,
                "columns": {"__sheet__": sheet.title, **cells},
            }


class WorkbookExtractor(Extractor):
    """One Bronze record per spreadsheet row."""

    target = "records"

    def handles(self, landed: Landed) -> bool:
        return landed.path.suffix.lower() in WORKBOOK_SUFFIXES

    def extract(self, landed: Landed) -> Iterator[dict[str, Any]]:
        try:
            workbook = _load(landed.path.read_bytes())
        except ExtractionError:
            raise
        except Exception as exc:  # noqa: BLE001 - surfaced with the path attached
            raise ExtractionError(str(landed.path), f"unreadable workbook: {exc}") from exc

        try:
            yield from _sheet_rows(workbook, landed.dataset or landed.path.stem)
        finally:
            workbook.close()


class ZippedWorkbookExtractor(Extractor):
    """Workbooks inside an archive.

    Bank Indonesia publishes its surveys as a ZIP holding a single workbook
    whose name carries the release month. The archive is landed rather than its
    contents, because the archive is what was published — so it is opened here
    instead.
    """

    target = "records"

    def handles(self, landed: Landed) -> bool:
        return landed.path.suffix.lower() in ARCHIVE_SUFFIXES

    def extract(self, landed: Landed) -> Iterator[dict[str, Any]]:
        try:
            archive = zipfile.ZipFile(io.BytesIO(landed.path.read_bytes()))
        except zipfile.BadZipFile as exc:
            raise ExtractionError(str(landed.path), f"unreadable archive: {exc}") from exc

        with archive:
            members = [
                name
                for name in archive.namelist()
                if name.lower().endswith(tuple(WORKBOOK_SUFFIXES))
                # Archives from macOS carry a shadow copy of every file.
                and not name.startswith("__MACOSX/")
            ]
            if not members:
                raise ExtractionError(str(landed.path), "archive holds no workbook")

            for member in members:
                workbook = _load(archive.read(member))
                try:
                    for row in _sheet_rows(workbook, landed.dataset or landed.path.stem):
                        # Which file inside the archive, since the name carries
                        # the release month and nothing else records it.
                        row["columns"]["__member__"] = member
                        yield row
                finally:
                    workbook.close()


def _spreadsheetml_rows(
    content: bytes, dataset: str, carried: dict[str, Any] | None = None
) -> Iterator[dict[str, Any]]:
    """Rows of one SpreadsheetML workbook, keyed by the header row.

    Keyed by header rather than by column letter, unlike the binary workbooks
    above: a portal serving this format is serving one table per response with
    its own header, not a formatted sheet with a title block above it.
    """
    try:
        root = ET.fromstring(content.decode("utf-8", errors="replace"))
    except ET.ParseError as exc:
        raise ExtractionError(dataset, f"unreadable SpreadsheetML: {exc}") from exc

    for sheet in root.iter(f"{{{SPREADSHEETML_NS}}}Worksheet"):
        name = sheet.get(f"{{{SPREADSHEETML_NS}}}Name") or ""
        header: list[str] | None = None
        number = 0
        for row in sheet.iter(f"{{{SPREADSHEETML_NS}}}Row"):
            cells = [(data.text or "").strip() for data in row.iter(f"{{{SPREADSHEETML_NS}}}Data")]
            if not any(cells):
                continue
            if header is None:
                header = cells
                continue

            number += 1
            if number > MAX_ROWS_PER_SHEET:
                break
            # `strict=False`: a row shorter than the header is a ragged table,
            # which these are, and the missing trailing cells are simply absent
            # rather than an error worth abandoning the file over.
            columns = {
                key: value
                for key, value in zip(header[:MAX_COLUMNS], cells[:MAX_COLUMNS], strict=False)
                if key
            }
            yield {
                "dataset": dataset,
                "row_number": number,
                "columns": {"__sheet__": name, **(carried or {}), **columns},
            }


class SpreadsheetMLExtractor(Extractor):
    """One Bronze record per row of a SpreadsheetML workbook.

    DJPK serves APBD realization this way — `application/xml` whose root is a
    `Workbook`. Claimed by inspecting the root element rather than by suffix
    alone, because `.xml` is also what an RSS feed and a sitemap are called, and
    handing one of those to this reader would yield nothing while stopping the
    text extractor from seeing it.
    """

    target = "records"

    def handles(self, landed: Landed) -> bool:
        if landed.path.suffix.lower() not in SPREADSHEETML_SUFFIXES:
            return False
        # Only the opening tag is needed, and these files run to megabytes.
        with landed.path.open("rb") as handle:
            head = handle.read(512)
        return b"urn:schemas-microsoft-com:office:spreadsheet" in head

    def extract(self, landed: Landed) -> Iterator[dict[str, Any]]:
        # What the source recorded about the request, carried onto every row.
        # DJPK's export names no fiscal year anywhere inside the file — the year
        # and month are what was asked for, not what was returned — so without
        # this the period is lost between RAW and Bronze and the rows cannot be
        # normalized at all. The file's own columns win a collision: what the
        # document says about itself outranks what we asked for.
        carried = {
            key: str(value)
            for key, value in (landed.extra or {}).items()
            if isinstance(value, str | int | float)
        }
        yield from _spreadsheetml_rows(
            landed.path.read_bytes(), landed.dataset or landed.path.stem, carried
        )
